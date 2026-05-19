from __future__ import annotations

import io
import re
from html import escape
from pathlib import Path

from docx import Document
from docx.shared import Pt

TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "templates"


def _load_template():
    path = TEMPLATE_DIR / "Mutual-NDA.md"
    if not path.exists():
        raise FileNotFoundError(f"MNDA template not found: {path}")
    return path.read_text()


def _fill_coverpage(data: dict) -> str:
    p1 = data.get("party1", {})
    p2 = data.get("party2", {})

    mnda_term_type = data.get("mnda_term_type", "expires")
    mnda_term_years = data.get("mnda_term_years")
    conf_term_type = data.get("confidentiality_term_type", "years")
    conf_years = data.get("confidentiality_years")

    if mnda_term_type == "expires" and mnda_term_years is not None:
        mnda_term_text = f"Expires {mnda_term_years} year(s) from Effective Date."
    else:
        mnda_term_text = "Continues until terminated in accordance with the terms of the MNDA."

    if conf_term_type == "years" and conf_years is not None:
        conf_term_text = f"{conf_years} year(s) from Effective Date, but in the case of trade secrets until Confidential Information is no longer considered a trade secret under applicable laws."
    else:
        conf_term_text = "In perpetuity."

    e = escape
    coverpage = f"""# Mutual Non-Disclosure Agreement

## Cover Page

**Purpose:** {e(data.get("purpose", ""))}

**Effective Date:** {e(data.get("effective_date", ""))}

**MNDA Term:** {e(mnda_term_text)}

**Term of Confidentiality:** {e(conf_term_text)}

**Governing Law:** {e(data.get("governing_law", ""))}

**Jurisdiction:** {e(data.get("jurisdiction", ""))}

---

## Parties

| | Party 1 | Party 2 |
|:---|:---:|:---:|
| **Name** | {e(p1.get("name", ""))} | {e(p2.get("name", ""))} |
| **Title** | {e(p1.get("title", ""))} | {e(p2.get("title", ""))} |
| **Company** | {e(p1.get("company", ""))} | {e(p2.get("company", ""))} |
| **Notice Address** | {e(p1.get("notice_address", ""))} | {e(p2.get("notice_address", ""))} |
| **Date** | {e(p1.get("date", ""))} | {e(p2.get("date", ""))} |
"""
    return coverpage


def _build_full_markdown(data: dict) -> str:
    coverpage = _fill_coverpage(data)
    mnda_template = _load_template()
    return f"{coverpage}\n\n---\n\n{mnda_template}"


def render_html(data: dict) -> str:
    import markdown
    md = _build_full_markdown(data)
    html_body = markdown.markdown(md, extensions=["tables", "fenced_code"])
    return f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{ font-family: Georgia, serif; max-width: 800px; margin: 40px auto; padding: 0 20px; line-height: 1.6; color: #333; }}
  h1 {{ font-size: 24px; border-bottom: 2px solid #333; padding-bottom: 8px; }}
  h2 {{ font-size: 18px; margin-top: 24px; }}
  table {{ border-collapse: collapse; width: 100%; margin: 16px 0; }}
  th, td {{ border: 1px solid #ccc; padding: 8px 12px; text-align: left; }}
  th {{ background: #f5f5f5; }}
  hr {{ border: none; border-top: 1px solid #ccc; margin: 24px 0; }}
  p {{ margin: 12px 0; }}
  strong {{ font-weight: bold; }}
</style>
</head>
<body>
{html_body}
</body>
</html>"""


def render_markdown(data: dict) -> str:
    return _build_full_markdown(data)


def render_docx(data: dict) -> bytes:
    md = _build_full_markdown(data)
    doc = Document()
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Georgia"
    font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    table_active = False
    current_table = None

    for line in md.split("\n"):
        stripped = line.strip()
        if not stripped:
            if table_active:
                table_active = False
            doc.add_paragraph("")
            continue

        if stripped.startswith("# "):
            if table_active:
                table_active = False
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            if table_active:
                table_active = False
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("---"):
            if table_active:
                table_active = False
            doc.add_paragraph("")
            continue
        elif stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.split("|")[1:-1]]
            if all(re.match(r"^:?-+:?$", c) for c in cells):
                continue
            if not table_active:
                table_active = True
                current_table = doc.add_table(rows=0, cols=len(cells))
                current_table.style = "Table Grid"

            row = current_table.add_row()
            is_header = len(current_table.rows) <= 2
            for i, cell_text in enumerate(cells):
                if i < len(row.cells):
                    row.cells[i].text = cell_text
                    for paragraph in row.cells[i].paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(10)
                            if is_header:
                                run.bold = True
        else:
            if table_active:
                table_active = False
            p = doc.add_paragraph()
            parts = stripped.split("**")
            for i, part in enumerate(parts):
                if not part:
                    continue
                run = p.add_run(part)
                run.font.size = Pt(11)
                if i % 2 == 1:
                    run.bold = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
